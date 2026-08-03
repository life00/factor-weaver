"""Parse financialdatadb xlsx sources into long-format parquet."""

from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd

from factor_weaver.data.tickers import normalize_ticker

_SECTIONS = (
    "Income Statement",
    "Balance Sheet",
    "Cash Flow Statement",
    "Non-GAAP Metrics",
    "Valuation Measures",
    "Valuation Ratios",
    "Liquidity/Efficiency Ratios",
    "Profitability Ratios",
    "Return Ratios",
)

_RATIO_SECTIONS = {
    "Valuation Ratios",
    "Profitability Ratios",
    "Return Ratios",
    "Liquidity/Efficiency Ratios",
}
_PER_SHARE_FIELDS = {"Earnings Per Share (Basic)", "Earnings Per Share (Diluted)"}

_COMPANY_COLUMNS = (
    "ticker",
    "company_name",
    "exchange",
    "country",
    "industry",
    "sector",
    "ipo_date",
    "description",
    "cik",
    "isin",
    "cusip",
)


def _paths(cfg: dict[str, Any]) -> tuple[Path, Path, Path]:
    """Resolve (raw_dir, fundamentals_out, companies_out) from the config."""
    c = cfg["financialdatadb"]
    return Path(c["raw_dir"]), Path(c["fundamentals_out"]), Path(c["companies_out"])


def _clean_cik(v) -> str | None:
    """Excel stores CIKs as integers, dropping leading zeros; pad to 10 digits."""
    if pd.isna(v):
        return None
    return f"{int(v):010d}"


def _parse_fundamentals_file(path: Path, sheet_names: list[str] | None = None) -> pd.DataFrame:
    """Extract long-format rows from one xlsx file (one sheet per ticker).

    Sheet layout: row 3 = report dates (col 2..N), rows 4+ = metric label in col 1,
    values in cols 2..N; rows whose label is a section header set the current section.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    records = []
    try:
        names = wb.sheetnames if sheet_names is None else sheet_names
        for sheet_name in names:
            ws = wb[sheet_name]
            it = ws.iter_rows(values_only=True)
            try:
                next(it)  # row 1: company name
                next(it)  # row 2: units note ('000s)
                dates = next(it)[1:]  # row 3: report dates
            except StopIteration:
                continue
            if not dates:
                continue
            section = None
            for row in it:
                label = row[0]
                if label is None:
                    break
                label = str(label)
                if label in _SECTIONS:
                    section = label
                    continue
                if section is None:
                    continue
                for date, value in zip(dates, row[1:]):
                    if date is not None and value is not None:
                        records.append((sheet_name, date, section, label, value))
    finally:
        wb.close()
    df = pd.DataFrame(records, columns=pd.Index(["ticker", "date", "section", "field", "value"]))
    # tickers with the dot format are legacy duplicates of dashed format tickers
    # those have to be dropped before normalization to avoid duplicate entries
    df = df.loc[~df["ticker"].str.contains(".", regex=False)]
    df["ticker"] = df["ticker"].map(normalize_ticker)
    df["value"] = pd.to_numeric(df["value"], errors="coerce")
    df["date"] = pd.to_datetime(df["date"], errors="coerce")
    mask = ~df["section"].isin(list(_RATIO_SECTIONS)) & ~df["field"].isin(list(_PER_SHARE_FIELDS))
    df.loc[mask, "value"] *= 1000.0
    df = df.dropna(subset=["value", "date"])
    # first-wins: resolves duplicate field names within a section (e.g. Gross Margin)
    return df.drop_duplicates(subset=["ticker", "date", "section", "field"], keep="first")


def parse_fundamentals(cfg: dict[str, Any]) -> None:
    """Walk financialdatadb *_tickers.xlsx files, melt wide→long per ticker sheet.

    Reads: <financialdatadb.raw_dir>/*_tickers.xlsx
    Writes: <financialdatadb.fundamentals_out>
    Columns: ticker, date, section, field, value
    Units: amounts in USD (scaled from source thousands); ratios and EPS unscaled.
    """
    raw_dir, out, _ = _paths(cfg)
    frames = []
    for path in sorted(raw_dir.glob("*_tickers.xlsx")):
        df = _parse_fundamentals_file(path)
        frames.append(df)
        print(f"  {path.name}: {len(df)} rows")
    out.parent.mkdir(parents=True, exist_ok=True)
    combined = pd.concat(frames, ignore_index=True)
    combined.to_parquet(out, index=False)
    print(f"wrote {out} ({len(combined)} rows)")


def _parse_companies_file(path: Path) -> pd.DataFrame:
    """Extract the static ticker→company mapping sheet (header on row 3)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        df = pd.DataFrame(ws.iter_rows(values_only=True)).iloc[3:, : len(_COMPANY_COLUMNS)]
    finally:
        wb.close()
    df.columns = pd.Index(_COMPANY_COLUMNS)
    df = df.replace("-", None).dropna(how="all")
    df["ticker"] = df["ticker"].map(normalize_ticker)
    df = df.drop_duplicates(subset=["ticker"], keep="first")
    df["ipo_date"] = pd.to_datetime(df["ipo_date"], errors="coerce")
    df["cik"] = df["cik"].apply(_clean_cik)
    for col in _COMPANY_COLUMNS:
        if col not in ("ipo_date", "cik"):
            df[col] = df[col].astype(str).where(df[col].notna(), None)
    return df


def parse_companies(cfg: dict[str, Any]) -> None:
    """Extract static ticker→company mapping (name, exchange, sector, CIK, ISIN, ...).

    Reads: <financialdatadb.raw_dir>/_US_listed_companies.xlsx
    Writes: <financialdatadb.companies_out>
    Columns: ticker, company_name, exchange, country, industry, sector,
             ipo_date, description, cik, isin, cusip
    """
    raw_dir, _, out = _paths(cfg)
    df = _parse_companies_file(raw_dir / "_US_listed_companies.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    df.to_parquet(out, index=False)
    print(f"wrote {out} ({len(df)} companies)")
